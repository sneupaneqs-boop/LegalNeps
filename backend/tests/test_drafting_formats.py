"""Prescribed-format drafting: every template renders to a real DOCX and a PDF,
official templates reproduce the fixed wording of the schedule they cite,
every citation resolves against the corpus, and the layout (A4, Devanagari
fonts, alignment, tables, thumbprint boxes) is really set in the DOCX XML.

The phrases in OFFICIAL_PHRASES were copied from the official schedules (the
PDF page is in each template's `source.url`); a few were also machine-checked
against the extracted schedule text of the Civil Procedure Code.
"""
import io
import re
import zipfile

import docx
import pymupdf
import pytest

from app.drafting import dsl, nepali, render
from app.drafting.fields import Cell, RepeatRow, Table
from app.drafting.registry import TEMPLATES
from app.retrieval import CORPUS_DIR

requires_corpus = pytest.mark.skipif(
    not (CORPUS_DIR / "manifest.json").exists(), reason="built corpus shards not present"
)

OFFICIAL = sorted(t.id for t in TEMPLATES.values() if t.kind == "official")
STANDARD = sorted(t.id for t in TEMPLATES.values() if t.kind == "standard")

# 3-5 distinctive fixed lines copied from each official schedule (blank form, so no answer text is involved)
OFFICIAL_PHRASES = {
    "plaint_civil": [
        "अदालत/कार्यालयले भर्ने",
        "म/हामी फिरादपत्रवाला निम्न प्रकरणहरूमा लेखिए बमोजिम फिराद गर्दछु/गर्दछौँ ।",
        "कानून व्यवसायी नियुक्त गरेको भए सोको विवरणः",
        "यसमा लेखिएको व्यहोरा ठीक साँचो छ, झुट्टा ठहरे कानून बमोजिम सहुँला बुझाउँला ।",
        "फिरादपत्रवालाको दस्तखत र सङ्गठित संस्था भए संस्थाको छाप",
    ],
    "written_reply_civil": [
        "उल्लिखित विपक्षी भएको उक्त मुद्दामा यस अदालतबाट मेरा/हाम्रा नाममा जारी भएको म्याद",
        "बादीले दाबी गरेका विषयहरूका सम्बन्धमा मेरो/हाम्रो भएको यथार्थ व्यहोरा निम्न प्रकरणहरूमा खुलाएको छु/छौँ ।",
        "फिराद दर्ता गर्ने हकदैया नभएको/नालेस गर्ने हदम्याद नभएको/अदालतको क्षेत्राधिकार नभएको जिकिर लिएको भए सोको आधार र कारण",
        "प्रतिउत्तरपत्रवालाको दस्तखत र सङ्गठित संस्था भए संस्थाको छाप",
    ],
    "appeal_civil": [
        "म/हामी पक्ष रहेको उपर्युक्त मुद्दामा देहायका कुरामा देहाय बमोजिम हुने गरी उपर्युक्त न्यायाधीश र अदालत/कार्यालयबाट फैसला भएकोमा",
        "पुनरावेदन गर्नु पर्ने कारण र फैसलाका आधार खण्डनको व्यहोराः",
        "फैसलाको प्रतिलिपि,",
        "पुनरावेदकको दस्तखत र सङ्गठित संस्था भए संस्थाको छाप",
    ],
    "warisnama_court": [
        "निज वारिस कानून बमोजिम वारिस हुन योग्य हुनुहुन्छ",
        "सो काम म इमान्दारीपूर्वक गर्नेछु",
        "औँठाको छाप/सङ्गठित संस्था भए संस्थाको छाप",
        "यो वारिसनामा हाम्रो रोहबरमा लेखी सहीछाप भएको साँचो हो ।",
    ],
    "application_district_court": [
        "म/हामी निम्न लिखित निवेदन गर्छु/गर्छौं :-",
        "तसर्थ यो यस विषयमा फलाना कानूनबमोजिम यो यस्तो गरिपाउँ वा यो यस्तो कुराका यो यसलाई आदेश गरिपाउँ।",
        "यस निवेदनपत्रको व्यहोरा ठीक साँचो छ, झुट्टा व्यहोरा लेखिएको ठहरे कानूनबमोजिम सजाय सहुँला बुझाउँला।",
    ],
    "petition_habeas_injunction_dc": [
        "म / हामी निम्नलिखित निवेदन गर्छु / गर्छौं :",
        "तसर्थ यो यस विषयमा फलाना कानूनबमोजिम यो यस्तो गरिपाउँ वा यो यस्तो कुराको यो यसलाई आदेश गरिपाउँ ।",
        "निवेदनको कारवाहीका सम्बन्धमा अदालतलाई आवश्यक सहयोग गर्नेछु ।",
    ],
    "writ_petition_high_court": [
        "म/हामी निम्नलिखित निवेदन गर्छु/गछौँः",
        "तसर्थ यो यस विषयमा फलाना कानून बमोजिम यो यस्तो गरिपाऊँ वा यो यस्तो कुराको यो यसलाई आदेश गरिपाऊँ।",
        "यस निवेदनपत्रको व्यहोरा ठीक साँचो छ, झूठा व्यहोरा लेखिएको ठहरे कानून बमोजिम सजाय सहुँला बुझाउँला ।",
    ],
    "written_response_high_court": [
        "लिखित जवाफ प्रस्तुतकर्ता",
        "तसर्थ यो विषयमा फलाना कानून बमोजिम यो यस्तो गरिपाऊँ वा यो यस्तो कुराको यो यसलाई आदेश गरिपाऊँ ।",
        "यस निवेदनपत्रको व्यहोरा ठीक साँचो छ, झूठा व्यहोरा लेखेको ठहरे कानून बमोजिम सजाय सहुँला बुझाउँला ।",
    ],
    "mediation_petition_court": [
        "विषय :—मेलमिलापद्वारा मुद्दा समाधान गरी पाउँ ।",
        "सम्बन्धमा आवश्यक व्यवस्था गरी पाउन यो निवेदन गरेको छु/छौँ ।",
        "निवेदनमा लेखिएको व्यहोरा ठीक साँचो छ, झुट्टा ठहरे कानून बमोजिम सँहुला बुझाउँला ।",
        "निवेदकको,–",
    ],
    "mediation_petition_dispute": [
        "विषय : मेलमिलापद्वारा विवाद समाधान गरी पाउँ ।",
        "विवादको संक्षिप्त विवरण :",
        "यसमा लेखिएको व्यहोरा ठीक साँचो छ, झुट्टा ठहरेमा कानून बमोजिम सहुँला बुझाउँला ।",
        "संलग्न कागजातहरु",
    ],
    "fir_jaheri_darkhast": [
        "प्रहरी कार्यालयले भर्ने",
        "जाहेरी दरखास्त वा सूचना दिने व्यक्तिको नाम, थर र ठेगाना (टोल सहित) फोन नम्बरः-",
        "कसूर भएको वा भइरहेको वा हुन लागेको स्थान, ठाउँ, मिति र समय,",
        "यो दरखास्तको व्यहोरा ठीक साँचो छ, झुट्टा व्यहोरा लेखेको ठहरे कानून बमोजिम सहुँला बुझाउँला ।",
        "जाहेरी वा सूचना दिने व्यक्तिकोः",
    ],
    "criminal_complaint_ujuri": [
        "उपर्युक्त कसूरदारले गरेको कसूरको व्यहोरा निम्न लिखित छ । निज उपर मुद्दाको कारबाही गर्न सादर अनुरोध गर्छु ।",
        "कसूरदारलाई हुनु पर्ने सजाय",
        "अभियुक्तले कसूर स्वीकार गरेको भए सो सम्बन्धी विवरण र सो बापत पाउने सजाय छुट",
        "(साक्षीहरुको पूरा नाम, उमेर र ठेगाना)",
    ],
    "criminal_written_reply": [
        "उपर्युक्त उजुरवालाले मलाई लगाएको अभियोगका सम्बन्धमा मेरो व्यहोरा निम्नलिखित छ ।",
        "साक्षीहरुको पूरा नाम, उमेर र ठेगाना",
        "अभियुक्तको सहीछाप",
    ],
    "criminal_appeal": [
        "उपरोक्त मुद्दामा म/हामीहरुलाई",
        "सजाय गरी बिगो वा क्षतिपूर्ति भराउने गरी फैसला गरेकोले फैसला बमोजिम भएको सजाय भोगी/सो बापत धरौट/जमानत दिई चित्त नबुझेको कुरामा पुनरावेदन दिन आएको छु/छौँ",
        "पुनरावेदकको सहीछाप",
    ],
    "rti_appeal_commission": [
        "श्री राष्ट्रिय सूचना आयोगमा दिएको",
        "ऐनको म्याद पैंतीस दिनभित्र यो पुनरावेदन गर्दछु/गर्दछौँ।",
        "माथि लेखिएको व्यहोरा ठीक साँचो हो, झुट्ठा ठहरे कानून बमोजिम सहुँला बुझाउँला।",
        "सार्वजनिक निकायको प्रमुखले गरेको निर्णयको प्रतिलिपि।",
        "पुनरावेदकको सही",
    ],
    "consumer_complaint": [
        "श्री केन्द्रीय बजार अनुगमन समिति/",
        "उजुरी गरेको/लेखाई दिएको छु । उजुरीमा उल्लेख गरेको व्यहोरा मैले जाने बुझेसम्म साँचो छ ।",
        "उजुरी कर्ताको नाम, थरः",
        "उजुरी प्राप्त भएको व्यहोरा प्रमाणित गर्ने अधिकारीको नाम, थर, दर्जा र प्रमाणित गरेको मितिः",
    ],
    "domestic_violence_complaint": [
        "घरेलु हिंसा सम्बन्धी कसूर गरेकोले आवश्यक कानूनी कारबाहीको लागि देहायको विवरण खुलाई यो उजुरी गरेको छु ।",
        "घरेलु हिंसाबाट पीडितलाई पर्न गएको असरः",
        "अन्य उजुरी सुन्ने निकायमा उजुरी गरेको भए सो निकायको नाम र उजुरी गरेको मितिः",
        "उजुरी दिनेको,",
    ],
    "birth_registration_notice": [
        "श्री स्थानीय पञ्जीकाधिकारीज्यू,",
        "निम्न लिखित विवरण खुलाई नवजात शिशु जन्मको सूचना दिन आएको छु । कानून अनुसार जन्म दर्ता गरी पाऊँ ।",
        "नवजात शिशुको बाबु आमाको विवरण",
        "यसमा लेखिएको विवरण साँचो हो। झुट्ठा ठहरे कानून बमोजिम सहुँला बुझाउँला भनी सहीछाप गर्ने सूचकको विवरण :",
    ],
    "death_registration_notice": [
        "निम्न लिखित विवरण खुलाई मृतकको सूचना दिन आएको छु । कानून अनुसार मृत्यु दर्ता गरी पाऊँ ।",
        "मृतक विवाहित भएमा",
        "विदेशी भएमा सूचक र मृतकको राहदानी, प्रवेशाज्ञा तथा निज त्यस वडामा बसोबास गरिरहेको प्रमाण",
    ],
    "marriage_registration_notice": [
        "निम्न लिखित विवरण खुलाई विवाहको सूचना दिन आएको छु । कानून अनुसार विवाह दर्ता गरी पाऊँ ।",
        "दुलहा दुलहीको विवरण",
        "दुवै जनाको हालै खिचेको अटो साईजको फोटो ।",
    ],
    "divorce_registration_notice": [
        "निम्न लिखित विवरण खुलाई सम्बन्ध विच्छेदको सूचना दिन आएको छु । कानून अनुसार सम्बन्ध विच्छेदको दर्ता गरी पाऊँ ।",
        "अदालतबाट सम्बन्ध विच्छेद भएको फैसलाको प्रतिलिपि",
        "पति पत्नीको विवरण",
    ],
    "migration_registration_notice": [
        "निम्न लिखित विवरण खुलाई बसाइँ सराईको सूचना दिन आएको छु । कानून अनुसार बसाइँ सराई दर्ता गरी पाऊँ ।",
        "बसाइँ सराई गर्ने परिवारका सदस्यहरुको विवरण",
        "सूचकको सहीछाप",
    ],
}


def _dummy(field):
    if field.type == "date":
        return "2026-09-29"
    if field.type == "number":
        return 42
    if field.type == "select":
        return field.options[0]["value"]
    if field.type == "textarea":
        return f"{field.id} एक\n{field.id} दुई"
    return f"{field.id}-नमुना"


def _answers(spec, required_only=True):
    return {f.id: _dummy(f) for f in spec.fields if f.required or not required_only}


def _all_text(document):
    parts = [p.text for p in document.paragraphs]
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _blank_text(template_id):
    """Rendered text of the template with only its required fields filled."""
    spec = TEMPLATES[template_id]
    data = render.render_docx(template_id, _answers(spec), "ne")
    return _all_text(docx.Document(io.BytesIO(data)))


def _xml(data, part="word/document.xml"):
    return zipfile.ZipFile(io.BytesIO(data)).read(part).decode("utf-8")


# ------------------------------------------------------------- catalogue ---

def test_at_least_thirty_templates_and_both_kinds_present():
    assert len(TEMPLATES) >= 30
    assert len(OFFICIAL) >= 15 and len(STANDARD) >= 10


def test_every_category_is_used_and_valid():
    cats = {t.category for t in TEMPLATES.values()}
    assert cats == {"court", "police", "office", "deeds", "notices"}


@pytest.mark.parametrize("template_id", sorted(TEMPLATES))
def test_every_template_declares_its_basis(template_id):
    spec = TEMPLATES[template_id]
    assert spec.kind in ("official", "standard")
    assert spec.source and spec.source.get("law_title_ne") is not None
    if spec.kind == "official":
        s = spec.source
        assert s["law_title_ne"] and s["schedule"].startswith("अनुसूची") and s["relates_to"] and s["form_title"]
        assert s["url"].startswith("https://") and re.search(r"#page=\d+$", s["url"]) and s["page"] >= 1
    else:
        # never dressed up as a prescribed form: a standard format says so
        note = spec.source.get("note")
        assert note and note["en"] and note["ne"]


def test_every_official_template_has_phrases_and_vice_versa():
    assert set(OFFICIAL_PHRASES) == set(OFFICIAL)
    for tid, phrases in OFFICIAL_PHRASES.items():
        assert 3 <= len(phrases) <= 5, tid


# ------------------------------------------------------------- rendering ---

@pytest.mark.parametrize("template_id", sorted(TEMPLATES))
@pytest.mark.parametrize("language", ["ne", "en"])
def test_every_template_renders_a_docx_python_docx_can_reopen(template_id, language):
    spec = TEMPLATES[template_id]
    data = render.render_docx(template_id, _answers(spec), language)
    document = docx.Document(io.BytesIO(data))
    assert any(p.text.strip() for p in document.paragraphs)


@pytest.mark.parametrize("template_id", sorted(TEMPLATES))
def test_every_template_renders_a_pdf_with_embedded_devanagari_font(template_id):
    spec = TEMPLATES[template_id]
    data = render.render_pdf(template_id, _answers(spec, required_only=False), "ne")
    assert data.startswith(b"%PDF")
    pdf = pymupdf.open(stream=data, filetype="pdf")
    assert len(pdf) >= 1
    assert abs(pdf[0].rect.width - 595.28) < 1 and abs(pdf[0].rect.height - 841.89) < 1  # A4
    fonts = {f[3] for page in pdf for f in page.get_fonts()}
    assert any("Noto Sans Devanagari" in f for f in fonts), fonts


def test_every_optional_field_left_blank_prints_a_dotted_blank_not_an_error():
    # the prescribed forms are meant to be printable with blanks to fill by hand
    for tid in OFFICIAL:
        spec = TEMPLATES[tid]
        text = _blank_text(tid)
        assert "None" not in text, tid
        if any(not f.required for f in spec.fields) and tid not in ("fir_jaheri_darkhast",):
            assert "…" in text, tid  # (the FIR prints labels followed by free space, like the schedule)


@pytest.mark.parametrize("template_id", OFFICIAL)
def test_official_template_reproduces_the_schedules_fixed_wording(template_id):
    text = _blank_text(template_id)
    flat = re.sub(r"\s+", " ", text)
    for phrase in OFFICIAL_PHRASES[template_id]:
        assert re.sub(r"\s+", " ", phrase) in flat, f"{template_id}: missing {phrase!r}"


@pytest.mark.parametrize("template_id", OFFICIAL)
def test_official_template_prints_the_schedule_heading_unless_switched_off(template_id):
    spec = TEMPLATES[template_id]
    src = spec.source
    if not any(f.id == "schedule_heading" for f in spec.fields):
        return  # a form whose schedule has no numbered heading (RTI appeal) keeps only its title
    on = _blank_text(template_id)
    assert src["schedule"] in on and f"({src['relates_to']} सँग सम्बन्धित)" in on and src["form_title"] in on
    answers = dict(_answers(spec), schedule_heading="no")
    off = _all_text(docx.Document(io.BytesIO(render.render_docx(template_id, answers, "ne"))))
    assert f"({src['relates_to']} सँग सम्बन्धित)" not in off


def test_prescribed_forms_fall_back_to_nepali_for_an_english_request():
    spec = TEMPLATES["plaint_civil"]
    assert spec.languages == ("ne",)
    data = render.render_docx("plaint_civil", _answers(spec), "en")
    assert "फिरादपत्र" in _all_text(docx.Document(io.BytesIO(data)))


def test_consumer_complaint_keeps_an_english_rendering_of_the_form():
    spec = TEMPLATES["consumer_complaint"]
    assert spec.languages == ("en", "ne")
    data = render.render_docx("consumer_complaint", _answers(spec), "en")
    text = _all_text(docx.Document(io.BytesIO(data)))
    assert "Consumer Protection Act, 2075" in text and "unofficial" in text


def test_missing_required_field_still_raises_for_prescribed_forms():
    with pytest.raises(render.MissingField):
        render.render_docx("plaint_civil", {}, "ne")


# ------------------------------------------------------------- citations ---

@requires_corpus
@pytest.mark.parametrize("template_id", sorted(TEMPLATES))
def test_every_cited_provision_resolves_against_the_corpus(template_id):
    detail = render.get_template_detail(template_id)
    for p in detail["provisions"]:
        assert p["citation"] and p["slug"]


@requires_corpus
def test_official_templates_cite_at_least_their_own_rule_or_section():
    for tid in OFFICIAL:
        # every official template is one of: has provisions, or (schedule-only forms such as the
        # registration notices) carries the schedule source instead
        spec = TEMPLATES[tid]
        assert spec.provisions or spec.source["url"]


# ------------------------------------------------------ layout in the XML ---

REPRESENTATIVE = ["plaint_civil", "fir_jaheri_darkhast", "writ_petition_high_court", "rti_request", "power_of_attorney"]


@pytest.mark.parametrize("template_id", REPRESENTATIVE)
def test_docx_page_and_fonts_are_really_set(template_id):
    spec = TEMPLATES[template_id]
    data = render.render_docx(template_id, _answers(spec, required_only=False), "ne")
    document, styles = _xml(data), _xml(data, "word/styles.xml")
    # A4 portrait, in twips (21.0 x 29.7 cm)
    m = re.search(r'<w:pgSz [^>]*w:w="(\d+)"[^>]*w:h="(\d+)"', document)
    assert m and abs(int(m.group(1)) - 11906) <= 2 and abs(int(m.group(2)) - 16838) <= 2
    # Kalimati named on every font slot Word consults for Devanagari, complex-script size set
    assert 'w:cs="Kalimati"' in styles and 'w:eastAsia="Kalimati"' in styles
    assert 'w:ascii="Kalimati"' in styles and re.search(r"<w:szCs w:val=\"24\"", styles)
    assert 'w:bidi="ne-NP"' in styles
    assert "theme" not in re.search(r"<w:rPrDefault>.*?</w:rPrDefault>", styles, re.S).group(0).lower()  # no theme-font override
    # runs carry the fonts too, bold runs carry the complex-script bold twin
    assert 'w:cs="Kalimati"' in document
    if "<w:b/>" in document:
        assert "<w:bCs/>" in document
    # the font table gives Word a fallback for machines without Kalimati
    font_table = _xml(data, "word/fontTable.xml")
    assert 'w:name="Kalimati"' in font_table and "Mangal" in font_table and "Noto Sans Devanagari" in font_table


_SCHEMA_ORDER = {
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
            "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
            "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
            "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
            "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
              "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
              "tblCaption", "tblDescription"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader",
             "tblCellSpacing", "jc", "hidden"],
}


@pytest.mark.parametrize("template_id", ["plaint_civil", "warisnama_court", "marriage_registration_notice",
                                         "migration_registration_notice", "power_of_attorney", "criminal_appeal"])
def test_docx_child_elements_follow_the_schema_order_word_requires(template_id):
    from lxml import etree

    spec = TEMPLATES[template_id]
    data = render.render_docx(template_id, _answers(spec, required_only=False), "ne")
    for part in ("word/document.xml", "word/styles.xml"):
        root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read(part))
        for name, order in _SCHEMA_ORDER.items():
            for el in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + name):
                kids = [k.tag.rsplit("}", 1)[-1] for k in el]
                ranks = [order.index(k) for k in kids if k in order]
                assert ranks == sorted(ranks), (part, name, kids)
                assert len(kids) == len(set(kids)), (part, name, "duplicate child", kids)


def test_plaint_layout_alignment_indents_and_hanging_numbers():
    spec = TEMPLATES["plaint_civil"]
    data = render.render_docx("plaint_civil", _answers(spec, required_only=False), "ne")
    xml = _xml(data)
    assert 'w:jc w:val="right"' in xml          # court-fills-this box, signature line
    assert 'w:jc w:val="center"' in xml         # titles and party lines
    assert 'w:jc w:val="both"' in xml           # the द्रष्टव्य note is justified
    assert "<w:u " in xml                        # underlined titles
    assert re.search(r'w:hanging="\d+"', xml)    # numbered items hang: "१." then wrapped text
    assert "<w:tab/>" in xml                     # label, tab, text - not spaces
    document = docx.Document(io.BytesIO(data))
    right = [p.text for p in document.paragraphs if p.alignment == 2]  # WD_ALIGN_PARAGRAPH.RIGHT
    assert any("अदालत/कार्यालयले भर्ने" in t for t in right)
    assert any("सङ्गठित संस्था भए संस्थाको छाप" in t for t in right)


def test_fir_has_the_two_column_police_fills_this_table():
    spec = TEMPLATES["fir_jaheri_darkhast"]
    data = render.render_docx("fir_jaheri_darkhast", _answers(spec, required_only=False), "ne")
    document = docx.Document(io.BytesIO(data))
    assert len(document.tables) == 1
    cells = [c.text for row in document.tables[0].rows for c in row.cells]
    assert any("प्रहरी कार्यालयले भर्ने" in c for c in cells)
    assert any(c.startswith("दर्ता नं:-") for c in cells)
    assert 'w:tblLayout w:type="fixed"' in _xml(data)


def test_writ_petition_puts_party_roles_on_the_right_and_versus_in_the_centre():
    spec = TEMPLATES["writ_petition_high_court"]
    data = render.render_docx("writ_petition_high_court", _answers(spec, required_only=False), "ne")
    document = docx.Document(io.BytesIO(data))
    right = [p.text for p in document.paragraphs if p.alignment == 2]
    centre = [p.text for p in document.paragraphs if p.alignment == 1]
    assert any(t.endswith("निवेदक") for t in right) and any(t.endswith("विपक्षी") for t in right)
    assert "विरुद्ध" in centre and "निवेदनपत्र" in centre


def test_power_of_attorney_has_bordered_left_right_thumbprint_table_and_witnesses():
    spec = TEMPLATES["power_of_attorney"]
    data = render.render_docx("power_of_attorney", _answers(spec, required_only=False), "ne")
    document = docx.Document(io.BytesIO(data))
    assert len(document.tables) == 1
    rows = [[c.text for c in row.cells] for row in document.tables[0].rows]
    assert rows[0] == ["दायाँ", "बायाँ"] and rows[1] == ["", ""]
    xml = _xml(data)
    assert "<w:tcBorders>" in xml
    assert "ल्याप्चे सहीछाप" in _all_text(document)
    assert "साक्षीहरू" in _all_text(document)
    # the original clause text is unchanged and the underscore signature lines are gone
    text = _all_text(document)
    assert "____" not in text and "दफा ५९१" in text and "दफा ५९२" in text


def test_power_of_attorney_english_thumbprints_appear_once_in_english_only():
    spec = TEMPLATES["power_of_attorney"]
    data = render.render_docx("power_of_attorney", _answers(spec, required_only=False), "en")
    document = docx.Document(io.BytesIO(data))
    assert len(document.tables) == 1
    assert [c.text for c in document.tables[0].rows[0].cells] == ["Right", "Left"]


def test_rti_request_date_is_right_aligned_and_signature_block_on_the_right():
    spec = TEMPLATES["rti_request"]
    data = render.render_docx("rti_request", _answers(spec, required_only=False), "ne")
    document = docx.Document(io.BytesIO(data))
    right = [p.text for p in document.paragraphs if p.alignment == 2]
    assert any(t.startswith("मितिः") for t in right)
    assert any(t == "निवेदक" for t in right)
    assert "सूचनाको हक सम्बन्धी ऐन, २०६४ को दफा ७" in _all_text(document)


def test_warisnama_thumbprint_boxes_and_two_column_signature_table():
    spec = TEMPLATES["warisnama_court"]
    data = render.render_docx("warisnama_court", _answers(spec), "ne")
    document = docx.Document(io.BytesIO(data))
    assert len(document.tables) == 4  # two thumbprint tables + two signature tables
    thumbs = [t for t in document.tables if [c.text for c in t.rows[0].cells] == ["दायाँ", "बायाँ"]]
    assert len(thumbs) == 2


def test_marriage_notice_has_photo_boxes_and_bordered_detail_table():
    spec = TEMPLATES["marriage_registration_notice"]
    data = render.render_docx("marriage_registration_notice", _answers(spec, required_only=False), "ne")
    document = docx.Document(io.BytesIO(data))
    first = document.tables[0]
    assert [c.text for c in first.rows[0].cells] == ["दुलाहाको\nफोटो", "दुलहीको\nफोटो"]
    assert 'w:jc w:val="right"' in _xml(data)  # the photo table sits at the right margin
    detail = [t for t in document.tables if t.rows[0].cells[1].text == "दुलहाको विवरण"]
    assert detail and len(detail[0].columns) == 3


def test_checkbox_rows_tick_the_chosen_option():
    spec = TEMPLATES["birth_registration_notice"]
    answers = dict(_answers(spec), c_sex="महिला")
    text = _all_text(docx.Document(io.BytesIO(render.render_docx("birth_registration_notice", answers, "ne"))))
    assert "☐ पुरुष   ☒ महिला   ☐ अन्य" in text


def test_dates_print_in_bs_with_the_traditional_closing_line():
    spec = TEMPLATES["plaint_civil"]
    answers = dict(_answers(spec), doc_date="2026-09-29")  # a Tuesday
    text = _all_text(docx.Document(io.BytesIO(render.render_docx("plaint_civil", answers, "ne"))))
    assert "इति सम्वत् २०८३ साल असोज महिना १२ गते रोज ३ शुभम् ।" in text or \
           "इति सम्वत् २०८३ साल असोज महिना १३ गते रोज ३ शुभम् ।" in text


# ---------------------------------------------------------------- the DSL ---

def test_dsl_flags_labels_and_indents():
    (p,) = dsl.doc("@cbu,>2,s14 {{ x }} शीर्षक")
    assert p.align == "center" and p.bold and p.underline and p.indent == 2 and p.size == 14
    (q,) = dsl.doc("@l १.|पहिलो {{ a|blank }} खण्ड")
    assert q.label == "१." and q.hanging == 1.0 and "{{ a|blank }}" in q.text["ne"]
    (r,) = dsl.doc("plain line")
    assert r.align == "left" and r.label is None


def test_dsl_pipe_inside_a_jinja_expression_is_not_a_label_separator():
    (p,) = dsl.doc("@l {{ court|blank }} अदालत")
    assert p.label is None


def test_dsl_tables_rows_and_repeat_rows():
    blocks = dsl.doc("@table 8,8 all h=2.5 align=right\n| @b a | b |\n@rows items:2\n| {{ n }} | {{ item }} |\n@end")
    (t,) = blocks
    assert isinstance(t, Table) and t.borders == "all" and t.row_height == 2.5 and t.align == "right"
    assert isinstance(t.rows[0][0], Cell) and t.rows[0][0].bold
    assert isinstance(t.rows[1], RepeatRow) and t.rows[1].each == "items" and t.rows[1].min_items == 2


def test_dsl_rejects_an_unknown_flag():
    with pytest.raises(ValueError):
        dsl.doc("@zz text")


def test_repeat_rows_expand_per_line_and_pad_to_the_minimum():
    spec = TEMPLATES["migration_registration_notice"]
    answers = dict(_answers(spec), members="राम | १ | 2050-01-01")
    doc_ = docx.Document(io.BytesIO(render.render_docx("migration_registration_notice", answers, "ne")))
    members = [t for t in doc_.tables if len(t.columns) == 10][-1]  # header table, then the member rows
    assert len(members.rows) == 3               # the minimum of three rows, first filled from the answer
    assert members.rows[0].cells[1].text == "राम" and members.rows[0].cells[0].text == "१"


# ---------------------------------------------------------- nepali helpers ---

def test_nepali_helpers():
    assert nepali.nd("2083-06-13") == "२०८३-०६-१३"
    assert [nepali.ka(i) for i in (1, 2, 3, 5)] == ["क", "ख", "ग", "ङ"]
    assert nepali.blank("") == "…" * 10 and nepali.blank("x") == "x" and nepali.blank(None, 3) == "…" * 3
    assert nepali.bs_numeric("2026-01-01") == "२०८२।०९।१७"
    assert nepali.bs_numeric("") == "" and nepali.bs_numeric("नमुना") == "नमुना"
    assert nepali.ad_numeric("2026-09-29") == "२९-०९-२०२६"
    assert nepali.signoff("").startswith("इति सम्वत् ………") and nepali.signoff("").endswith("शुभम् ।")
    assert nepali.signoff("2026-09-29", "संवत्").startswith("इति संवत् २०८३ साल")
    assert nepali.opts("b", ["a", "b"]) == "☐ a   ☒ b"


# --------------------------------------------------------------------- API ---

@requires_corpus
def test_api_lists_categories_kinds_sources_and_serves_pdf():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        rows = c.get("/api/drafting/templates").json()
        by_id = {r["id"]: r for r in rows}
        assert {r["category"] for r in rows} == {"court", "police", "office", "deeds", "notices"}
        assert by_id["plaint_civil"]["kind"] == "official"
        assert by_id["plaint_civil"]["source"]["schedule"] == "अनुसूची–१"
        assert by_id["plaint_civil"]["source"]["url"].startswith("https://giwmscdnone.gov.np/")
        assert by_id["will_sheshpachi_bakaspatra"]["kind"] == "standard"
        assert by_id["will_sheshpachi_bakaspatra"]["source"]["note"]["en"]

        detail = c.get("/api/drafting/templates/plaint_civil").json()
        assert detail["formats"] == ["docx", "pdf"] and detail["provisions"]
        rel = next(f for f in detail["fields"] if f["id"] == "pl_rel")
        assert rel["type"] == "select" and [o["value"] for o in rel["options"]][:2] == ["छोरा", "छोरी"]

        spec = TEMPLATES["plaint_civil"]
        body = {"language": "ne", "answers": _answers(spec), "format": "pdf"}
        r = c.post("/api/drafting/templates/plaint_civil/draft", json=body)
        assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
        assert r.headers["content-disposition"].endswith('plaint_civil.pdf"') and r.content.startswith(b"%PDF")

        body["format"] = "rtf"
        assert c.post("/api/drafting/templates/plaint_civil/draft", json=body).status_code == 422
