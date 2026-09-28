"""S9: drafting engine - questionnaire answers -> rendered DOCX, bilingual.
Table-driven per STRATEGY's "done when" bar for this session: generated
DOCX opens, snapshot-style content checks per template.

Citation-bearing templates are checked against the real corpus (same
requires_corpus pattern as tests/test_playbooks.py and
tests/test_calculators.py) via app.drafting.render.get_template_detail,
which resolves every provision through app.playbooks.resolve_provision.
"""
import io

import docx
import pytest

from app.drafting import render
from app.drafting.registry import TEMPLATES
from app.retrieval import CORPUS_DIR

requires_corpus = pytest.mark.skipif(
    not (CORPUS_DIR / "manifest.json").exists(), reason="built corpus shards not present"
)

EXPECTED_TEMPLATE_IDS = {
    "legal_notice_salary",
    "legal_notice_deposit",
    "rental_agreement",
    "power_of_attorney",
    "affidavit",
    "consumer_complaint",
}

VALID_ANSWERS = {
    "legal_notice_salary": dict(
        sender_name="राम बहादुर श्रेष्ठ", sender_address="काठमाडौं",
        recipient_name="ABC प्रा.लि.", recipient_address="ललितपुर",
        job_title="लेखापाल", unpaid_period="जेठ-असार २०८२", unpaid_amount=45000, demand_days=15,
    ),
    "legal_notice_deposit": dict(
        sender_name="सीता देवी", sender_address="भक्तपुर",
        recipient_name="हरि प्रसाद", recipient_address="काठमाडौं",
        deposit_amount=20000, vacate_date="2026-01-01", demand_days=15,
    ),
    "rental_agreement": dict(
        landlord_name="हरि प्रसाद", landlord_address="काठमाडौं", landlord_citizenship_no="12-34-56",
        tenant_name="सीता देवी", tenant_address="भक्तपुर", tenant_citizenship_no="65-43-21",
        house_location="काठमाडौं-५", kitta_no="234", purpose="बसोबास",
        start_date="2026-01-01", duration_months=12, monthly_rent=15000,
        payment_terms="प्रत्येक महिनाको सुरुमा", utility_responsibility="बहालवाला आफैं",
        sublet_allowed="होइन",
    ),
    "power_of_attorney": dict(
        principal_name="राम बहादुर श्रेष्ठ", principal_address="काठमाडौं", principal_citizenship_no="11-22-33",
        agent_name="गीता श्रेष्ठ", agent_address="ललितपुर", agent_citizenship_no="44-55-66",
        purpose_description="मेरो तर्फबाट जग्गा रजिस्ट्रेसन प्रक्रिया गर्ने।",
    ),
    "affidavit": dict(
        declarant_name="विनोद कुमार", declarant_address="पोखरा", declarant_citizenship_no="77-88-99",
        purpose_of_affidavit="प्रमाणपत्रमा नाम फरक परेको स्पष्ट गर्न",
        statement_text="मेरो नागरिकतामा 'विनोद' र शैक्षिक प्रमाणपत्रमा 'बिनोद' लेखिएको छ, दुवै एउटै व्यक्ति हुँ।",
    ),
    "consumer_complaint": dict(
        complainant_name="अनिता राई", complainant_address="विराटनगर", complainant_phone="9800000000",
        seller_name="XYZ इलेक्ट्रोनिक्स", seller_address="विराटनगर",
        product_or_service="मोबाइल फोन", purchase_date="2026-01-01", amount_involved=25000,
        complaint_details="किनेको एक हप्तामै फोन बन्द भयो, पसलले साट्न वा फिर्ता दिन मानेन।",
        relief_sought="रकम फिर्ता वा नयाँ फोनसाटी दिनुहोस्।",
    ),
}


def test_registry_has_exactly_the_six_s9_templates():
    assert set(TEMPLATES) == EXPECTED_TEMPLATE_IDS


def test_list_templates_matches_registry():
    ids = {t["id"] for t in render.list_templates()}
    assert ids == EXPECTED_TEMPLATE_IDS


@pytest.mark.parametrize("template_id", sorted(EXPECTED_TEMPLATE_IDS))
@pytest.mark.parametrize("language", ["ne", "en"])
def test_generated_docx_opens(template_id, language):
    data = render.render_docx(template_id, VALID_ANSWERS[template_id], language)
    document = docx.Document(io.BytesIO(data))
    assert len(document.paragraphs) > 0
    assert any(p.text.strip() for p in document.paragraphs)


def test_legal_notice_salary_includes_answers_and_citation_ne():
    data = render.render_docx("legal_notice_salary", VALID_ANSWERS["legal_notice_salary"], "ne")
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert "राम बहादुर श्रेष्ठ" in text
    assert "ABC प्रा.लि." in text
    assert "रु. 45000" in text
    assert "दफा ३५" in text
    assert "दफा १६२" in text


def test_rental_agreement_includes_all_mandatory_clauses_ne():
    data = render.render_docx("rental_agreement", VALID_ANSWERS["rental_agreement"], "ne")
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    for expected in ["हरि प्रसाद", "सीता देवी", "काठमाडौं-५", "234", "15000", "दफा ३८६"]:
        assert expected in text


def test_power_of_attorney_optional_valid_until_field():
    answers_without = VALID_ANSWERS["power_of_attorney"]
    data = render.render_docx("power_of_attorney", answers_without, "ne")
    text_without = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert "मान्य रहनेछ" not in text_without

    answers_with = dict(answers_without, valid_until="2027-01-01")
    data = render.render_docx("power_of_attorney", answers_with, "ne")
    text_with = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert "2027-01-01" in text_with
    assert "मान्य रहनेछ" in text_with


def test_missing_required_field_raises():
    incomplete = dict(VALID_ANSWERS["legal_notice_salary"])
    del incomplete["unpaid_amount"]
    with pytest.raises(render.MissingField):
        render.render_docx("legal_notice_salary", incomplete, "ne")


def test_unknown_template_raises():
    with pytest.raises(render.UnknownTemplate):
        render.render_docx("not_a_real_template", {}, "ne")


def test_invalid_language_raises():
    with pytest.raises(ValueError):
        render.render_docx("affidavit", VALID_ANSWERS["affidavit"], "fr")


@requires_corpus
@pytest.mark.parametrize("template_id", sorted(EXPECTED_TEMPLATE_IDS))
def test_template_detail_provisions_resolve_against_corpus(template_id):
    detail = render.get_template_detail(template_id)
    for provision in detail["provisions"]:
        assert provision["citation"]
        assert provision["slug"]


def test_affidavit_has_no_citation_by_design():
    # A generic sworn statement isn't itself created under one dedicated
    # statute in the corpus, unlike the other 5 templates - documented, not
    # an oversight.
    assert TEMPLATES["affidavit"].provisions == []


# ------------------------------------------------------------------- API ---

@requires_corpus
def test_drafting_api_endpoints():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        r = c.get("/api/drafting/templates")
        assert r.status_code == 200
        assert {t["id"] for t in r.json()} == EXPECTED_TEMPLATE_IDS

        r = c.get("/api/drafting/templates/legal_notice_salary")
        assert r.status_code == 200
        body = r.json()
        assert len(body["fields"]) == len(VALID_ANSWERS["legal_notice_salary"])
        assert body["provisions"]

        assert c.get("/api/drafting/templates/does-not-exist").status_code == 404

        r = c.post(
            "/api/drafting/templates/legal_notice_salary/draft",
            json={"language": "ne", "answers": VALID_ANSWERS["legal_notice_salary"]},
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        document = docx.Document(io.BytesIO(r.content))
        assert len(document.paragraphs) > 0

        r = c.post(
            "/api/drafting/templates/legal_notice_salary/draft",
            json={"language": "ne", "answers": {}},
        )
        assert r.status_code == 400

        assert c.post(
            "/api/drafting/templates/does-not-exist/draft",
            json={"language": "ne", "answers": {}},
        ).status_code == 404
