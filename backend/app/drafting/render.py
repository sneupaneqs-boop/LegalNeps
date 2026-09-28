"""Questionnaire answers -> rendered DOCX (S9).

Each template's paragraphs are Jinja source strings (see registry.py); this
module fills in the questionnaire answers plus a couple of always-available
context values (today's date in BS, via S8's date calculator) and lays the
result out as a real .docx with python-docx.
"""
from __future__ import annotations

import datetime
import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Environment, StrictUndefined
from jinja2.exceptions import TemplateError, UndefinedError

from ..calculators.dates import ad_to_bs
from ..playbooks import resolve_provision
from .fields import TemplateSpec
from .registry import TEMPLATES

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

_JINJA_ENV = Environment(undefined=StrictUndefined, autoescape=False)


class UnknownTemplate(ValueError):
    pass


class MissingField(ValueError):
    pass


def list_templates() -> list[dict]:
    return [{"id": t.id, "title": t.title, "description": t.description} for t in TEMPLATES.values()]


def get_template(template_id: str) -> TemplateSpec:
    spec = TEMPLATES.get(template_id)
    if spec is None:
        raise UnknownTemplate(f"{template_id!r} is not a known drafting template")
    return spec


def get_template_detail(template_id: str) -> dict:
    spec = get_template(template_id)
    return {
        "id": spec.id,
        "title": spec.title,
        "description": spec.description,
        "fields": [
            {"id": f.id, "label": f.label, "type": f.type, "required": f.required, "help": f.help}
            for f in spec.fields
        ],
        "provisions": [resolve_provision(p) for p in spec.provisions],
    }


def _build_context(spec: TemplateSpec, answers: dict) -> dict:
    ctx: dict = {}
    for f in spec.fields:
        value = answers.get(f.id, f.default)
        if f.required and (value is None or value == ""):
            raise MissingField(f"{f.id!r} is required for template {spec.id!r}")
        ctx[f.id] = value if value is not None else ""
    today = datetime.date.today()
    ctx["today_ad"] = today.isoformat()
    ctx["today_bs"] = str(ad_to_bs(today))
    return ctx


def render_paragraphs(template_id: str, answers: dict, language: str) -> list[tuple[str, bool, str]]:
    """[(text, bold, align), ...] for every non-empty paragraph of
    `template_id` rendered in `language` ("en" or "ne")."""
    if language not in ("en", "ne"):
        raise ValueError("language must be 'en' or 'ne'")
    spec = get_template(template_id)
    ctx = _build_context(spec, answers)
    out: list[tuple[str, bool, str]] = []
    for p in spec.paragraphs:
        source = p.text.get(language)
        if not source:
            continue
        try:
            rendered = _JINJA_ENV.from_string(source).render(**ctx)
        except UndefinedError as exc:
            raise MissingField(str(exc)) from exc
        except TemplateError as exc:  # pragma: no cover - malformed template source, not user input
            raise ValueError(f"template {template_id!r} paragraph failed to render: {exc}") from exc
        rendered = rendered.strip()
        if rendered:
            out.append((rendered, p.bold, p.align))
    return out


def render_docx(template_id: str, answers: dict, language: str = "ne") -> bytes:
    paragraphs = render_paragraphs(template_id, answers, language)
    document = docx.Document()
    for text, bold, align in paragraphs:
        para = document.add_paragraph()
        para.alignment = _ALIGN[align]
        for i, line in enumerate(text.split("\n")):
            if i > 0:
                para.add_run().add_break()
            run = para.add_run(line)
            run.bold = bold
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
