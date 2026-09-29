"""Resolved layout blocks -> a real .docx (python-docx).

Page and type set-up follows Nepali court / office practice: A4, ~2.5 cm
margins, a Unicode Devanagari face named on *every* font slot Word consults
(`w:ascii`, `w:hAnsi`, `w:cs`, `w:eastAsia`) with a font-table fallback, and the
complex-script size / bold / italic twins (`w:szCs`, `w:bCs`, `w:iCs`) that Word
uses for Devanagari runs. Alignment, indents, hanging labels, page breaks and
tables are written as real Word properties (never spaces or tabs faked in
the text).
"""
from __future__ import annotations

import io
import re

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .fields import TemplateSpec
from .inline import split_inline

# First choice is the font Nepal's government and courts type in; the
# fontTable's <w:altName> tells Word which installed face to use instead when
# Kalimati is not on the machine (Mangal ships with Windows / Office; Noto
# Sans Devanagari is the free cross-platform choice).
PRIMARY_FONT = "Kalimati"
FALLBACK_FONTS = ("Noto Sans Devanagari", "Mangal")

PAGE_W_CM, PAGE_H_CM = 21.0, 29.7  # A4
MARGIN_CM = {"top": 2.5, "bottom": 2.5, "left": 2.7, "right": 2.3}
LINE_SPACING = 1.15
DEFAULT_SPACE_AFTER = 6.0

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_TABLE_ALIGN = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}


# Word validates the order of child elements against the WordprocessingML schema, so every
# element added by hand goes in at its schema position (python-docx's own setters already do).
_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1]


def _insert_ordered(parent, child, order: list) -> None:
    rank = order.index(_local(child))
    for existing in parent:
        name = _local(existing)
        if name in order and order.index(name) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def _get_or_add(parent, tag: str, order: list):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        _insert_ordered(parent, el, order)
    return el


def _set_fonts(rpr, name: str = PRIMARY_FONT) -> None:
    rfonts = _get_or_add(rpr, "w:rFonts", _RPR_ORDER)
    for attr in list(rfonts.attrib):  # drop theme font indirections that override names
        if "theme" in attr.lower():
            del rfonts.attrib[attr]
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(slot), name)


def _set_size(rpr, pt: float) -> None:
    for tag in ("w:sz", "w:szCs"):
        _get_or_add(rpr, tag, _RPR_ORDER).set(qn("w:val"), str(int(round(pt * 2))))


def _set_flag(rpr, tag: str, on: bool) -> None:
    el = rpr.find(qn(tag))
    if on:
        if el is None:
            _insert_ordered(rpr, OxmlElement(tag), _RPR_ORDER)
    elif el is not None:
        rpr.remove(el)


def _style_defaults(document, size_pt: float) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.size = Pt(size_pt)
    rpr = normal.element.get_or_add_rPr()
    _set_fonts(rpr)
    _set_size(rpr, size_pt)
    pf = normal.paragraph_format
    pf.space_after = Pt(DEFAULT_SPACE_AFTER)
    pf.space_before = Pt(0)
    pf.line_spacing = LINE_SPACING
    pf.widow_control = True
    # document-wide defaults (docDefaults) too, so styles other than Normal (table text) match
    root = styles.element
    defaults = root.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr_default is not None:
            _set_fonts(rpr_default)
            _set_size(rpr_default, size_pt)
            lang = _get_or_add(rpr_default, "w:lang", _RPR_ORDER)
            lang.set(qn("w:val"), "en-US")
            lang.set(qn("w:bidi"), "ne-NP")
            lang.set(qn("w:eastAsia"), "en-US")


def _page_setup(document) -> None:
    sec = document.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(PAGE_W_CM)
    sec.page_height = Cm(PAGE_H_CM)
    sec.top_margin = Cm(MARGIN_CM["top"])
    sec.bottom_margin = Cm(MARGIN_CM["bottom"])
    sec.left_margin = Cm(MARGIN_CM["left"])
    sec.right_margin = Cm(MARGIN_CM["right"])


def _add_font_table_fallback(document) -> None:
    """Register Kalimati in word/fontTable.xml with an altName so Word/LibreOffice
    substitute a Devanagari font that is actually installed."""
    for part in document.part.package.iter_parts():
        if str(part.partname) != "/word/fontTable.xml":
            continue
        blob = part.blob.decode("utf-8")
        alt = ",".join(FALLBACK_FONTS)
        entry = (
            f'<w:font w:name="{PRIMARY_FONT}"><w:altName w:val="{alt}"/>'
            '<w:charset w:val="01"/><w:family w:val="auto"/><w:pitch w:val="variable"/></w:font>'
        )
        if f'w:name="{PRIMARY_FONT}"' not in blob:
            blob = re.sub(r"</w:fonts>\s*$", entry + "</w:fonts>", blob)
            part._blob = blob.encode("utf-8")
        return


def _run(par, text: str, bold=False, italic=False, underline=False, size: float | None = None) -> None:
    run = par.add_run(text)
    rpr = run._r.get_or_add_rPr()
    _set_fonts(rpr)
    if bold:
        run.bold = True
        _set_flag(rpr, "w:bCs", True)
    if italic:
        run.italic = True
        _set_flag(rpr, "w:iCs", True)
    if underline:
        run.underline = True
    if size:
        _set_size(rpr, size)


def _write_text(par, text: str, bold, italic, underline, size) -> None:
    """Text with \\n as real line breaks inside one paragraph."""
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            par.add_run().add_break()
        for seg, b, u in split_inline(line, bold, underline):
            _run(par, seg, b, italic, u, size)


def _paragraph(document_or_cell, p, default_size: float, in_cell: bool = False):
    par = document_or_cell.add_paragraph()
    pf = par.paragraph_format
    par.alignment = _ALIGN[p.align]
    if getattr(p, "indent", 0) or getattr(p, "hanging", 0):
        left = (p.indent or 0) + (p.hanging or 0)
        pf.left_indent = Cm(left)
        if p.hanging:
            pf.first_line_indent = Cm(-p.hanging)
            pf.tab_stops.add_tab_stop(Cm(left), WD_TAB_ALIGNMENT.LEFT)
    if p.space_before is not None:
        pf.space_before = Pt(p.space_before)
    if p.space_after is not None:
        pf.space_after = Pt(p.space_after)
    if p.keep_next:
        pf.keep_with_next = True
    if p.page_break_before:
        pf.page_break_before = True
    size = p.size if p.size else None
    if p.label:
        _run(par, p.label, p.bold, p.italic, p.underline, size)
        par.add_run("\t")
    _write_text(par, p.text, p.bold, p.italic, p.underline, size)
    return par


def _set_cell_borders(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    _insert_ordered(tcpr, borders, _TCPR_ORDER)


def _table(document, t, default_size: float) -> None:
    ncols = len(t.widths)
    table = document.add_table(rows=len(t.rows), cols=ncols)
    table.alignment = _TABLE_ALIGN[t.align]
    table.autofit = False
    # (`autofit = False` above already writes <w:tblLayout w:type="fixed"/> in schema order)
    for r_idx, row in enumerate(t.rows):
        tr = table.rows[r_idx]
        if t.row_height and (r_idx == len(t.rows) - 1 or len(t.rows) == 1):
            tr.height = Cm(t.row_height)
            tr.height_rule = WD_ROW_HEIGHT.AT_LEAST
        for c_idx in range(ncols):
            cell = tr.cells[c_idx]
            cell.width = Cm(t.widths[c_idx])
            if t.borders == "all":
                _set_cell_borders(cell)
            spec = row[c_idx] if c_idx < len(row) else None
            first = cell.paragraphs[0]
            if spec is None:
                continue
            first.alignment = _ALIGN[spec.align]
            first.paragraph_format.space_after = Pt(2)
            first.paragraph_format.space_before = Pt(2)
            _write_text(first, spec.text, spec.bold, spec.italic, spec.underline, t.size)
            if spec.valign == "center":
                from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # a small spacer so the next paragraph does not butt against the table
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(t.space_after if t.space_after is not None else 2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)


def build_docx(spec: TemplateSpec, blocks: list) -> bytes:
    from .render import RBreak, RPara, RTable

    document = docx.Document()
    _page_setup(document)
    _style_defaults(document, spec.font_size)
    _add_font_table_fallback(document)
    document.core_properties.title = spec.title.get("ne") or spec.title.get("en") or spec.id
    document.core_properties.author = "Kanooni Sathi"
    document.core_properties.language = "ne-NP"
    for b in blocks:
        if isinstance(b, RPara):
            _paragraph(document, b, spec.font_size)
        elif isinstance(b, RTable):
            _table(document, b, spec.font_size)
        elif isinstance(b, RBreak):
            document.add_page_break()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
