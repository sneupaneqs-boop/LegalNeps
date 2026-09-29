"""Render an audit result as a DOCX report (python-docx), in English or Nepali."""
from __future__ import annotations

import datetime
import io

import docx
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .audit import DISCLAIMER, STATUS_ORDER

_LABELS = {
    "title": {"en": "Contract audit report", "ne": "सम्झौता परीक्षण प्रतिवेदन"},
    "file": {"en": "File", "ne": "फाइल"},
    "type": {"en": "Contract type", "ne": "सम्झौताको प्रकार"},
    "date": {"en": "Date", "ne": "मिति"},
    "clauses": {"en": "Clauses read", "ne": "पढिएका दफा/बुँदा"},
    "summary": {"en": "Summary", "ne": "सारांश"},
    "finding": {"en": "Finding", "ne": "नतिजा"},
    "where": {"en": "Where in the contract", "ne": "सम्झौताको कहाँ"},
    "todo": {"en": "What to do", "ne": "के गर्ने"},
    "basis": {"en": "Legal basis", "ne": "कानूनी आधार"},
    "not_found": {"en": "Not found in the contract", "ne": "सम्झौतामा भेटिएन"},
    "best_practice": {"en": "Best practice (the law leaves this to the parties)",
                      "ne": "उत्तम अभ्यास (कानूनले पक्षहरूमै छाडेको)"},
    "statutory": {"en": "Required or limited by law", "ne": "कानूनले तोकेको वा सीमित गरेको"},
    "truncated": {"en": "Note: the contract was long and only its first part was analysed.",
                  "ne": "टिप्पणी: सम्झौता लामो भएकाले सुरुको भाग मात्र विश्लेषण गरियो।"},
    "no_findings": {"en": "No findings in this group.", "ne": "यो समूहमा कुनै नतिजा छैन।"},
}
_STATUS = {
    "issue": {"en": "Issues", "ne": "समस्या"},
    "warning": {"en": "Warnings", "ne": "चेतावनी"},
    "missing": {"en": "Missing", "ne": "नभएको"},
    "info": {"en": "Suggestions", "ne": "सुझाव"},
    "ok": {"en": "Looks fine", "ne": "ठीक देखिन्छ"},
}
_COLOURS = {"issue": "B42318", "warning": "B54708", "missing": "B54708", "info": "175CD3", "ok": "067647"}


def _fonts(doc: docx.document.Document) -> None:
    """Latin text in Calibri, Devanagari (complex script) in Mangal so Word
    renders Nepali correctly even on a machine without our fonts."""
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), "Mangal")


def _cell(cell, text: str, *, bold: bool = False, colour: str | None = None, size: float = 9) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour)


def _add(cell, text: str, *, bold: bool = False, italic: bool = False, size: float = 9) -> None:
    para = cell.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)


def render_report(audit: dict, language: str = "en") -> bytes:
    """`audit` is the dict returned by audit.run_audit (or the API's JSON)."""
    lang = language if language in ("en", "ne") else "en"
    L = lambda key: _LABELS[key][lang]  # noqa: E731
    document = docx.Document()
    _fonts(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.8))

    document.add_heading(L("title"), level=1)
    ctype = (audit.get("contract_type_title") or {}).get(lang) or audit.get("contract_type", "")
    meta = [
        (L("file"), audit.get("filename") or "-"),
        (L("type"), ctype),
        (L("date"), datetime.date.today().isoformat()),
        (L("clauses"), str(audit.get("clause_count", 0))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Light List"
    for k, v in meta:
        row = table.add_row().cells
        _cell(row[0], k, bold=True)
        _cell(row[1], v)

    note = document.add_paragraph()
    run = note.add_run(DISCLAIMER[lang])
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("B42318")
    if audit.get("truncated"):
        document.add_paragraph(L("truncated"))

    document.add_heading(L("summary"), level=2)
    counts = audit.get("summary") or {}
    st = document.add_table(rows=1, cols=len(STATUS_ORDER))
    st.style = "Light Grid"
    for i, status in enumerate(STATUS_ORDER):
        _cell(st.rows[0].cells[i], f"{_STATUS[status][lang]}: {counts.get(status, 0)}", bold=True,
              colour=_COLOURS[status])

    findings = audit.get("findings") or []
    for status in STATUS_ORDER:
        group = [f for f in findings if f.get("status") == status]
        if not group:
            continue
        document.add_heading(f"{_STATUS[status][lang]} ({len(group)})", level=2)
        tbl = document.add_table(rows=1, cols=4)
        tbl.style = "Light Grid"
        for i, key in enumerate(("finding", "where", "todo", "basis")):
            _cell(tbl.rows[0].cells[i], L(key), bold=True)
        for f in group:
            cells = tbl.add_row().cells
            title = (f.get("title") or {}).get(lang) or (f.get("title") or {}).get("en", "")
            _cell(cells[0], title, bold=True, colour=_COLOURS.get(status))
            if f.get("facts"):
                _add(cells[0], "; ".join(f"{k} = {v:g}" if isinstance(v, float) else f"{k} = {v}"
                                         for k, v in f["facts"].items()), italic=True, size=8)
            if f.get("clause_label"):
                _cell(cells[1], f["clause_label"], bold=True)
                if f.get("quote"):
                    _add(cells[1], f"“{f['quote']}”", italic=True, size=8.5)
            else:
                _cell(cells[1], L("not_found") if f.get("not_stated") else "-")
            rec = f.get("recommendation") or {}
            _cell(cells[2], rec.get(lang) or rec.get("en", "") if status != "ok" else "")
            prov = f.get("provision") or {}
            citation = (prov.get("citation") if lang == "ne" else prov.get("citation_en")) or prov.get("citation", "")
            _cell(cells[3], citation)
            if prov.get("url"):
                _add(cells[3], prov["url"], size=7.5)
            _add(cells[3], L(f.get("basis", "statutory")), italic=True, size=8)

    footer = document.add_paragraph()
    footer.paragraph_format.space_before = Pt(12)
    frun = footer.add_run(DISCLAIMER[lang])
    frun.italic = True
    frun.font.size = Pt(8.5)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
